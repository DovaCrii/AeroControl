"""LV-355 F1: la tubería de exportación (`apps/reporting/export.py`).

Qué fija, por orden de gravedad:

1. **Neutralización de fórmulas en toda celda de texto**, también en el pie de
   totales, el preámbulo y los encabezados (con un valor centinela: un `=` al
   principio es lo que Excel ejecuta al abrir).
2. **Un solo escritor por formato**: ningún informe vuelve a escribir su propio
   `Workbook` ni su propio CSV con BOM (guardián de fuentes).
3. Nombres de archivo `stem_AAAA-MM-DD`.
4. El filtro de faena compartido (S6) trata lo malformado como «sin filtro».
"""

import csv
import re
from datetime import date
from io import BytesIO, StringIO
from pathlib import Path

import pytest
from django.conf import settings

from apps.reporting.export import (
    BOM,
    Table,
    cost_center_filter,
    dated_filename,
    download,
    to_csv,
    to_docx,
    to_xlsx,
)

SENTINEL = '=HYPERLINK("http://evil","x")'
DANGEROUS = ("=", "+", "-", "@")


def _table(**kwargs):
    defaults = dict(
        headers=("Name", "Count"),
        rows=((SENTINEL, 3), ("+piloto", 4), ("@cuenta", 5), ("-menos", 6)),
        sheet="Datos",
        heading="Titulo",
        footer=("TOTAL", 18),
    )
    defaults.update(kwargs)
    return Table(**defaults)


def _csv_cells(content):
    text = content.decode("utf-8")
    assert text.startswith(BOM), "falta el BOM: Excel no detectaría UTF-8"
    return list(csv.reader(StringIO(text[1:])))


def _xlsx_cells(content):
    from openpyxl import load_workbook

    workbook = load_workbook(BytesIO(content))
    return [
        cell.value
        for sheet in workbook.worksheets
        for row in sheet.iter_rows()
        for cell in row
    ]


class TestFormulasAreInertEverywhere:
    def test_a_csv_cell_is_neutralised_in_rows_footer_preamble_and_headers(self):
        table = _table(
            headers=("=cmd", "ok"), footer=(SENTINEL, 1), heading="=encabezado"
        )

        cells = [
            cell
            for row in _csv_cells(to_csv([table], preamble=[["@titulo"], ["-corte"]]))
            for cell in row
        ]

        assert not [c for c in cells if c.startswith(DANGEROUS)], cells
        assert "'=cmd" in cells and "'@titulo" in cells and "'-corte" in cells
        assert "'" + SENTINEL in cells

    def test_an_xlsx_cell_is_neutralised_in_rows_footer_preamble_and_headers(self):
        table = _table(headers=("=cmd", "ok"), footer=(SENTINEL, 1))

        values = _xlsx_cells(to_xlsx([table], preamble=["@titulo", "-corte"]))

        text = [v for v in values if isinstance(v, str)]
        assert not [v for v in text if v.startswith(DANGEROUS)], text
        assert "'=cmd" in text and "'@titulo" in text and "'-corte" in text
        assert "'" + SENTINEL in text

    def test_a_negative_number_in_the_footer_is_a_number_not_text(self):
        """`-5` como número no es una fórmula: sólo el **texto** se neutraliza."""
        from openpyxl import load_workbook

        content = to_xlsx([_table(footer=("TOTAL", -5))])
        sheet = load_workbook(BytesIO(content)).active

        assert sheet.cell(row=6, column=2).value == -5

    def test_the_docx_writes_text_as_is(self):
        """Un párrafo de Word no ejecuta nada: neutralizar ahí ensuciaría el
        texto con un apóstrofo que nadie pidió (idéntico a antes)."""
        from docx import Document

        content = to_docx("T", [("table", _table())])

        document = Document(BytesIO(content))
        assert document.tables[0].rows[1].cells[0].text == SENTINEL


class TestTheXlsxLayout:
    def test_without_a_preamble_the_header_is_row_one(self):
        from openpyxl import load_workbook

        sheet = load_workbook(BytesIO(to_xlsx([_table()]))).active

        assert sheet.title == "Datos"
        assert sheet.cell(row=1, column=1).value == "Name"
        assert sheet.freeze_panes == "A2"
        assert sheet.auto_filter.ref == "A1:B6"  # el pie va dentro del filtro
        assert sheet.cell(row=1, column=1).font.bold
        assert sheet.cell(row=6, column=1).font.bold

    def test_with_a_preamble_the_filter_is_anchored_on_the_header(self):
        from openpyxl import load_workbook

        sheet = load_workbook(
            BytesIO(to_xlsx([_table(footer=None)], preamble=["Corte", "Otra"]))
        ).active

        assert sheet.cell(row=1, column=1).value == "Corte"
        assert sheet.cell(row=4, column=1).value == "Name"  # una fila en blanco
        assert sheet.auto_filter.ref == "A4:B8"
        assert sheet.freeze_panes == "A5"

    def test_the_width_follows_the_content_with_a_ceiling(self):
        from openpyxl import load_workbook

        sheet = load_workbook(
            BytesIO(to_xlsx([_table(rows=(("x" * 200, 1),), footer=None)]))
        ).active

        assert sheet.column_dimensions["A"].width == 40
        assert sheet.column_dimensions["B"].width == len("Count") + 2

    def test_one_sheet_per_table_in_order(self):
        from openpyxl import load_workbook

        workbook = load_workbook(
            BytesIO(to_xlsx([_table(sheet="Uno"), _table(sheet="Dos", footer=None)]))
        )

        assert workbook.sheetnames == ["Uno", "Dos"]

    def test_an_empty_table_still_has_its_header_and_a_valid_filter(self):
        from openpyxl import load_workbook

        sheet = load_workbook(BytesIO(to_xlsx([_table(rows=(), footer=None)]))).active

        assert sheet.auto_filter.ref == "A1:B1"


class TestTheCsvLayout:
    def test_the_layout_is_preamble_blank_headings_headers_rows_and_footer(self):
        rows = _csv_cells(
            to_csv(
                [_table(heading="Uno"), _table(heading="Dos", footer=None)],
                preamble=[["Reporte"], ["Corte", "2026-10-26"]],
                trailing_blank=True,
            )
        )

        assert rows[:3] == [["Reporte"], ["Corte", "2026-10-26"], []]
        assert rows[3] == ["Uno"] and rows[4] == ["Name", "Count"]
        assert rows[9] == ["TOTAL", "18"]
        assert rows[10] == [] and rows[11] == ["Dos"]
        assert rows[-1] == []

    def test_the_line_ending_is_crlf(self):
        assert b"\r\n" in to_csv([_table()])


class TestTheFilenames:
    def test_it_is_stem_underscore_date(self):
        assert (
            dated_filename("aerocontrol-catastro", date(2026, 10, 26), "pdf")
            == "aerocontrol-catastro_2026-10-26.pdf"
        )

    def test_the_download_names_the_attachment(self):
        response = download(b"x", "text/csv", "a_2026-10-26.csv")

        assert (
            response["Content-Disposition"] == 'attachment; filename="a_2026-10-26.csv"'
        )


@pytest.mark.django_db
class TestTheCostCentreFilter:
    def test_absent_malformed_and_unknown_values_are_no_filter(self):
        assert cost_center_filter({}) is None
        assert cost_center_filter({"cost_center": ""}) is None
        assert cost_center_filter({"cost_center": "not-a-uuid"}) is None
        assert (
            cost_center_filter({"cost_center": "3f2b8c1e-1111-4222-8333-444455556666"})
            is None
        )

    def test_an_active_cost_centre_is_found_and_an_archived_one_is_not(self):
        from apps.registry.models import CostCenter

        live = CostCenter.objects.create(code="F1", name="Viva")
        archived = CostCenter.objects.create(
            code="F2", name="Archivada", is_active=False
        )

        assert cost_center_filter({"cost_center": str(live.pk)}) == live
        assert cost_center_filter({"cost_center": str(archived.pk)}) is None


class TestThereIsOneWriterPerFormat:
    """El defecto que esto cierra: cuatro `Workbook()` y cuatro CSV con BOM,
    copiados, cada uno con su propia (y distinta) neutralización."""

    REPORT_SOURCES = [
        "apps/compliance/report_views.py",
        "apps/registry/report_views.py",
        "apps/workboard/views.py",
    ]

    @pytest.mark.parametrize("source", REPORT_SOURCES)
    def test_a_report_view_does_not_write_its_own_files(self, source):
        text = (Path(settings.BASE_DIR) / source).read_text(encoding="utf-8")

        assert "Workbook(" not in text, "XLSX escrito a mano: usar export.to_xlsx"
        assert "csv.writer" not in text, "CSV escrito a mano: usar export.to_csv"
        assert not re.search("ufeff|﻿", text), "BOM a mano: usar export.to_csv"
        assert "from docx" not in text, "DOCX escrito a mano: usar export.to_docx"
