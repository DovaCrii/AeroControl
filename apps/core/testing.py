"""Ayudantes compartidos por la suite de tests (T4.1).

`login_as` estaba copiado **33 veces**, 24 de ellas byte a byte, y las nueve
variantes sólo agregaban grupos, pertenencia a un tenant, o devolvían también el
usuario. Cada copia es una oportunidad de que el montaje de permisos de un test
diverja del de los demás sin que nadie lo note -- y el montaje de permisos es
justo lo que estos tests están afirmando.

No es un `conftest.py` con fixtures, que era lo que proponía la fila del plan.
Un ayudante importable deduplica lo mismo **sin tocar la firma de cada función de
test**: adoptarlo cuesta cambiar dos líneas por archivo en vez de trescientas
firmas, y una migración grande de tests que se hace a mano es exactamente donde
se cuela el error que la suite ya no puede cazar, porque la suite es lo que se
está moviendo. Las fixtures de datos (`two_tenant_world`) siguen pendientes y se
harán cuando haya un segundo lector real que las pida.
"""

import re

from django.contrib.auth.models import Group, Permission, User
from django.test import Client


def login_as(*codenames, groups=(), member_of=None):
    """Un cliente autenticado con **exactamente** estos permisos.

    El nombre de usuario se deriva de los permisos, así que dos clientes con
    permisos distintos no chocan en la restricción de unicidad y el que falla se
    identifica solo en el error.

    `groups` agrega grupos por nombre (el destinatario de las notificaciones,
    por ejemplo). `member_of` agrega el usuario a un `OperationalTenant`, que es
    como `visible_tenant_ids` resuelve el alcance -- una M2M `members`, no un
    campo del perfil.

    El usuario queda en `client.user`: los tests que lo necesitaban devolvían
    una tupla, y una tupla obliga a desempaquetar en todos los que no lo
    necesitan.
    """
    # `nosec B106`: es la contraseña de un usuario que sólo existe dentro de una
    # base de datos de test que se destruye al terminar. Bandit no marcaba esto
    # cuando el ayudante vivía copiado en archivos `test_*.py` (excluidos del
    # análisis); al extraerlo a un módulo normal, sí — y la respuesta correcta es
    # justificar la excepción acá, no ampliar la exclusión a todo el módulo.
    user = User.objects.create_user(  # nosec B106
        f"u-{'-'.join(codenames) or 'none'}", password="pw"
    )
    if codenames:
        user.user_permissions.add(*Permission.objects.filter(codename__in=codenames))
    for name in groups:
        user.groups.add(Group.objects.get_or_create(name=name)[0])
    if member_of is not None:
        member_of.members.add(user)
    client = Client()
    assert client.login(username=user.username, password="pw")  # nosec B106
    client.user = user
    return client


def pin_today_mid_month(monkeypatch):
    """Fija `timezone.localdate()` al **día 15 del mes en curso** y lo devuelve.

    `LV-257`: cuatro tests del informe afirmaban que el mes en curso "sigue
    abierto" usando el día real, y **fallaban todos los últimos días de mes** —
    el código considera cerrado el mes su último día, a propósito (lo fija
    `test_the_last_day_of_the_month_is_already_closed`). Se encontró corriendo la
    suite con el reloj movido al 31 de diciembre; el 1 de octubre y el 1 de enero
    pasaban. Es la lección de `LV-223` en `AGENTS.md`: la fecha es andamio y hay
    que declararla.

    Día 15 del mes **real**, no una fecha absoluta: los registros siguen
    sellándose con la hora real (`auto_now_add`), y una fecha de otro mes los
    dejaría fuera del período que el test mira.

    Con argumento, `localdate(valor)` sigue convirtiendo de verdad: sólo cambia
    la pregunta "¿qué día es hoy?".
    """
    from django.utils import timezone

    real = timezone.localdate
    pinned = real().replace(day=15)

    def localdate(value=None, timezone=None):
        if value is None:
            return pinned
        return real(value, timezone)

    monkeypatch.setattr("django.utils.timezone.localdate", localdate)
    return pinned


def pin_now(monkeypatch, moment):
    """Fija `timezone.now()` al instante `moment` (aware) y lo devuelve.

    `pin_today_mid_month` sólo fija la fecha; lo que depende de la HORA (la ventana del centelleo
    GNSS, `LV-352`) necesita fijar el instante. `localdate()`/`localtime()` sin argumento parten
    de `now()`, así que siguen coherentes con este valor.
    """
    from django.utils import timezone

    assert timezone.is_aware(moment), "pin_now exige un datetime con zona horaria"
    monkeypatch.setattr("django.utils.timezone.now", lambda: moment)
    return moment


def workbook_grid(content):
    """Un XLSX como datos comparables: hojas, celdas, negritas, filtro y anchos.

    `LV-355` F1: la prueba de «contenido idéntico antes/después» de la tubería de
    exportación. Lo que importa de un libro es lo que alguien ve al abrirlo --
    las hojas en su orden, el valor de cada celda, qué filas van en negrita,
    dónde está el filtro y el panel congelado, y el ancho de cada columna --, y
    todo eso cabe en un diccionario que se guarda como JSON y se compara.
    """
    from io import BytesIO

    from openpyxl import load_workbook

    workbook = load_workbook(BytesIO(content))
    sheets = {}
    for sheet in workbook.worksheets:
        sheets[sheet.title] = {
            "rows": [
                [cell.value for cell in row] for row in sheet.iter_rows(min_row=1)
            ],
            "bold_rows": [
                row[0].row for row in sheet.iter_rows() if row and row[0].font.bold
            ],
            "freeze": sheet.freeze_panes,
            "filter": sheet.auto_filter.ref,
            "widths": {
                letter: dimension.width
                for letter, dimension in sorted(sheet.column_dimensions.items())
            },
        }
    return {"sheets": sheets, "order": workbook.sheetnames}


def docx_grid(content):
    """Un DOCX como datos: párrafos con su estilo, y cada tabla fila por fila."""
    from io import BytesIO

    from docx import Document

    document = Document(BytesIO(content))
    return {
        "paragraphs": [[p.style.name, p.text] for p in document.paragraphs],
        "tables": [
            [[cell.text for cell in row.cells] for row in table.rows]
            for table in document.tables
        ],
        "left_margins": [section.left_margin.inches for section in document.sections],
    }


def assert_golden(path, data):
    """Compara `data` con el JSON guardado en `path`.

    Con `UPDATE_SNAPSHOTS=1` lo reescribe (y hay que revisar el diff). Los
    archivos `golden` de `LV-355` F1 se generaron **con el código anterior a la
    tubería**: que sigan pasando es la prueba de que unificar los escritores no
    movió una sola celda.
    """
    import json
    import os

    text = json.dumps(data, indent=1, ensure_ascii=False, sort_keys=True) + "\n"
    if os.environ.get("UPDATE_SNAPSHOTS") == "1":
        path.parent.mkdir(parents=True, exist_ok=True)
        path.write_text(text, encoding="utf-8", newline="\n")
        return
    assert path.exists(), f"falta {path.name}: UPDATE_SNAPSHOTS=1 para crearlo"
    assert text == path.read_text(encoding="utf-8"), (
        f"la salida cambió respecto de {path.name}; si es a propósito, "
        "UPDATE_SNAPSHOTS=1 y revisar el diff"
    )


_TEMPLATE_COMMENT = re.compile(
    r"\{%\s*comment\s*%\}.*?\{%\s*endcomment\s*%\}", re.DOTALL
)


def without_template_comments(text):
    """Los `{% comment %}` de una plantilla, en blanco — **conservando líneas**.

    Todo guardián que lee plantillas como texto lo necesita, porque un comentario
    **no se renderiza**: exigirle `scope` a un `<th>` escrito ahí adentro, o
    prohibir la palabra `<table>` porque aparece explicando por qué la tabla se
    dibuja como se dibuja, es pedirle algo a marcado que no existe en la página.

    Estaba escrito **tres veces** cuando se extrajo —en el guardián de
    traducciones (`LV-169`), en el de `scope` y en el de `UX-07`— y las tres
    veces por el mismo tropiezo: alguien documenta una decisión nombrando la
    etiqueta de la que habla, y el guardián lo trata como código. Tres copias del
    mismo recorte es cómo una de ellas se queda sin arreglar.

    Se reemplaza por espacios y no se recorta, para que los números de línea que
    los guardianes reportan sigan apuntando al lugar real del archivo.
    """
    return _TEMPLATE_COMMENT.sub(
        lambda match: re.sub(r"[^\n]", " ", match.group(0)), text
    )
