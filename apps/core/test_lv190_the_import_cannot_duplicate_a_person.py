"""LV-190: el import del Capítulo 1 ya no puede duplicar una persona.

Encontrado el 2026-08-31 al preparar el cruce que el usuario pidió — los
operadores inscritos en el Rev 17 contra los que tiene AeroControl, para subir
los faltantes, como se hizo con las aeronaves. El manual trae **48 fichas** y
producción tiene 42, así que la corrida iba a crear unas seis. Ahí estaba el
problema.

`partition()` cruzaba los operadores **sólo por `employee_id`**, mientras las
aeronaves van por matrícula **y** número de serie; el docstring de `LV-134` dice
que esa segunda llave "es la parte que evita el duplicado real". El RUT —que
`LV-143` declaró la llave natural— no participaba. Y tres piezas se alineaban
mal:

1. Una ficha creada a mano antes de `LV-169` lleva un `employee_id` que no es
   `RUT-…`, así que el import no la reconocía.
2. `apply_report` crea con `Operator.objects.create()`, que **no llama a
   `clean()`**, donde vive la única comprobación de RUT repetido.
3. `Operator.rut` es `CharField(blank=True)` **sin `unique=True`**: la base
   tampoco lo frenaba.

La aeronave duplicada revienta contra un índice. La persona duplicada entraba en
silencio — lo que `LV-143` vino a cerrar, por la puerta de atrás, en el padrón
que la DGAC espeja.

La segunda mitad, ya sembrada: como `create()` no pasa por `clean()`, las fichas
que este comando creó tienen el RUT **con puntos**, tal como viene del manual, y
`operator_with_rut` compara contra la forma canónica. Por eso el cruce normaliza
los dos lados, y por eso el import pasa a guardar el RUT normalizado.
"""

import json

import pytest
from django.core.management import call_command
from django.core.management.base import CommandError
from docx import Document as DocxDocument

from apps.registry.models import Operator

RUT = "17.816.266-7"
CANONICAL = "17816266-7"
EMPLOYEE_ID = "RUT-178162667"


@pytest.fixture
def source(tmp_path):
    """Un Capítulo 1 mínimo con la forma de la Rev 17: una ficha, una nave."""
    document = DocxDocument()
    document.add_paragraph("1.5-\tDOTACIÓN OPERADOR RPA")
    document.add_paragraph("\t1)\tPERMANENTES")
    document.add_paragraph("\t\t\tNOMBRE\t\t:\tCristobal Muñoz Montiel")
    document.add_paragraph(f"RUT\t\t:\t{RUT}")
    document.add_paragraph("Credencial N°\t:\t8172")
    document.add_paragraph("Tipo\t\t:\tOperador RPA")
    document.add_paragraph("Habilitaciones\t:\tMatrice Series")
    document.add_paragraph("Dirección\t:\tDiagonal Santa Elena 2605")
    document.add_paragraph("Teléfono\t:\t987282880")
    document.add_paragraph("Email\t\t:\tcmunoz@jej.cl")
    document.add_paragraph("2) EVENTUALES (NO APLICA)")
    service_table = document.add_table(rows=2, cols=3)
    service_table.rows[0].cells[0].text = "AERONAVES"
    service_table.rows[1].cells[2].text = "Fotografía"
    inventory = document.add_table(rows=2, cols=8)
    inventory.rows[0].cells[0].text = "Propietario"
    for index, value in enumerate(
        [
            "J.E.J. Ingeniería S.A",
            "DJI / MATRICE 4 ENTERPRISE",
            "SER-7213",
            "RPA-7213",
            "1.420",
            "1.420",
            "VLOS",
            "NO",
        ]
    ):
        inventory.rows[1].cells[index].text = value
    path = tmp_path / "chapter1_rev17.docx"
    document.save(path)
    return path


def _report(source, capsys, *args):
    call_command("chapter1_docx_import", "--source", str(source), "--json", *args)
    return json.loads(capsys.readouterr().out)


@pytest.mark.django_db
class TestTheDuplicateThatCouldGetThrough:
    def test_a_hand_made_fiche_with_the_same_rut_is_a_conflict(self, source, capsys):
        """El caso real: una ficha dada de alta a mano antes de `LV-169`, con su
        propio `employee_id`. El import no la reconocía y creaba la persona de
        nuevo — sin que la base pudiera impedirlo."""
        Operator.objects.create(
            employee_id="P-01", full_name="Cristóbal Muñoz", rut=CANONICAL
        )

        report = _report(source, capsys)

        assert report["conflicts"]
        assert "P-01" in report["conflicts"][0]

    def test_apply_refuses_the_conflict_even_with_skip_existing(self, source):
        """Un conflicto detiene la corrida **siempre**, con bandera o sin ella:
        no es "ya está", es "el manual y la base no coinciden", y adivinar cuál
        manda sobre una persona es cómo se crea la ficha fantasma."""
        Operator.objects.create(
            employee_id="P-01", full_name="Cristóbal Muñoz", rut=CANONICAL
        )

        with pytest.raises(CommandError, match="disagree"):
            call_command(
                "chapter1_docx_import",
                "--source",
                str(source),
                "--apply",
                "--skip-existing",
            )
        assert Operator.objects.count() == 1

    def test_the_rut_is_matched_normalized_on_both_sides(self, source, capsys):
        """No es precaución teórica: las fichas que este comando ya creó tienen el
        RUT **con puntos**, porque `create()` no pasa por `clean()`. Comparar en
        crudo no habría encontrado justo a las personas que él cargó."""
        Operator.objects.create(
            employee_id="P-01", full_name="Cristóbal Muñoz", rut=RUT
        )

        assert _report(source, capsys)["conflicts"]


@pytest.mark.django_db
class TestWhatMustKeepWorking:
    def test_a_fiche_this_import_created_is_skipped_not_duplicated(
        self, source, capsys
    ):
        """La corrida normal: el ID coincide y el RUT también. Es la misma
        persona y se salta, que es lo que `LV-134` dejó funcionando."""
        Operator.objects.create(
            employee_id=EMPLOYEE_ID, full_name="Cristobal Muñoz Montiel", rut=CANONICAL
        )

        report = _report(source, capsys)

        assert report["conflicts"] == []
        assert f"operator:{EMPLOYEE_ID}" in report["skipped"]

    def test_an_old_fiche_with_no_rut_on_file_is_still_the_same_person(
        self, source, capsys
    ):
        """El ID coincide y la ficha no tiene el RUT escrito. Tratarlo como
        conflicto habría convertido en un bloqueo cada ficha vieja incompleta."""
        Operator.objects.create(
            employee_id=EMPLOYEE_ID, full_name="Cristobal Muñoz Montiel", rut=""
        )

        report = _report(source, capsys)

        assert report["conflicts"] == []
        assert f"operator:{EMPLOYEE_ID}" in report["skipped"]

    def test_a_genuinely_new_person_is_still_created(self, source):
        """El contrapeso: la llave nueva no puede bloquear lo que hay que cargar,
        que es el punto de todo el ejercicio."""
        call_command("chapter1_docx_import", "--source", str(source), "--apply")

        assert Operator.objects.filter(employee_id=EMPLOYEE_ID).exists()

    def test_it_stores_the_rut_normalized(self, source):
        """Y deja de sembrar valores que la app no puede encontrar: con el RUT
        con puntos, `operator_with_rut` no ve a esta persona y el formulario de
        alta dejaría crear su duplicado a mano."""
        call_command("chapter1_docx_import", "--source", str(source), "--apply")

        assert Operator.objects.get(employee_id=EMPLOYEE_ID).rut == CANONICAL
