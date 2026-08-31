"""LV-195: el import compara el número de serie normalizado.

Encontrado el 2026-08-31 al preparar el cruce del Capítulo 1 Rev 17 contra
producción, y encontrado **en los datos reales**: el manual trae dos seriales
partidos por un espacio —`RPA-4401` (`1581F5FHC245700 D181D`) y `RPA-4436`
(`1581F5FHC245800 DWTY6`)—, casi con seguridad un salto de línea dentro de la
celda del Word.

`partition()` comparaba `row["serial_number"]` crudo contra los seriales de la
base, que `Aircraft.save()` guarda pasados por `normalize_serial` (mayúsculas y
sin espacios, ADR-0002 §2). O sea: comparaba contra una forma que la base nunca
tiene. Las dos aeronaves salían como *"matrícula ya existe, serie nueva"*, y
**un conflicto detiene la corrida completa** con o sin `--skip-existing` — así
que un espacio en el documento bloqueaba la carga entera, con un mensaje que
acusaba al registro DGAC de una discrepancia que no existía.

Es el mismo defecto de forma que `LV-190`: una llave comparada sin la función
canónica que la app ya tiene y usa en todos los demás caminos (`Aircraft.save()`,
`Battery.save()`, la sincronización con AeroLink).

El serial se sigue **extrayendo tal como viene** del manual: el informe describe
la fuente, y ver el espacio es lo que delata el problema del documento. Lo que se
normaliza es la comparación — y al crear, `save()` lo normaliza igual que
siempre.
"""

import json

import pytest
from django.core.management import call_command
from docx import Document as DocxDocument

from apps.registry.models import Aircraft, normalize_serial

# El caso real del Rev 17, con el espacio incluido.
SERIAL_DEL_MANUAL = "1581F5FHC245700 D181D"
SERIAL_EN_LA_BASE = "1581F5FHC245700D181D"


def _source(tmp_path, *, registration="RPA-4401", serial=SERIAL_DEL_MANUAL):
    document = DocxDocument()
    document.add_paragraph("1.5-\tDOTACIÓN OPERADOR RPA")
    document.add_paragraph("\t1)\tPERMANENTES")
    document.add_paragraph("\t\t\tNOMBRE\t\t:\tCristobal Muñoz Montiel")
    document.add_paragraph("RUT\t\t:\t17.816.266-7")
    document.add_paragraph("Credencial N°\t:\t8172")
    document.add_paragraph("Tipo\t\t:\tOperador RPA")
    document.add_paragraph("Habilitaciones\t:\tMatrice Series")
    document.add_paragraph("Dirección\t:\tDiagonal Santa Elena 2605")
    document.add_paragraph("Teléfono\t:\t987282880")
    document.add_paragraph("Email\t\t:\tcmunoz@jej.cl")
    document.add_paragraph("2) EVENTUALES (NO APLICA)")
    service_table = document.add_table(rows=2, cols=3)
    service_table.rows[0].cells[0].text = "AERONAVES"
    service_table.rows[1].cells[2].text = "Fotografía"
    inventory = document.add_table(rows=2, cols=8)
    inventory.rows[0].cells[0].text = "Propietario"
    for index, value in enumerate(
        [
            "J.E.J. Ingeniería S.A",
            "DJI / MAVIC 3 ENTERPRISE",
            serial,
            registration,
            "1.420",
            "1.420",
            "VLOS",
            "NO",
        ]
    ):
        inventory.rows[1].cells[index].text = value
    path = tmp_path / "chapter1_rev17.docx"
    document.save(path)
    return path


def _aircraft(registration, serial):
    return Aircraft.objects.create(
        registration=registration,
        type="RPA",
        model="MAVIC 3 ENTERPRISE",
        manufacturer="DJI",
        serial_number=serial,
    )


def _report(source, capsys):
    call_command("chapter1_docx_import", "--source", str(source), "--json")
    return json.loads(capsys.readouterr().out)


@pytest.mark.django_db
class TestTheSpaceInTheWordCell:
    def test_the_same_aircraft_is_skipped_and_not_a_conflict(self, tmp_path, capsys):
        """El caso real de `RPA-4401`: la aeronave ya está, y el manual escribe su
        serie con un espacio en medio. Es la misma máquina."""
        _aircraft("RPA-4401", SERIAL_EN_LA_BASE)

        report = _report(_source(tmp_path), capsys)

        assert report["conflicts"] == []
        assert "aircraft:RPA-4401" in report["skipped"]

    def test_a_space_no_longer_blocks_the_whole_run(self, tmp_path):
        """Lo que costaba de verdad: el conflicto detiene la corrida entera, así
        que un espacio en el Word impedía cargar **todo** lo demás — incluidas las
        fichas de personal que no tienen nada que ver con esa celda."""
        _aircraft("RPA-4401", SERIAL_EN_LA_BASE)

        call_command(
            "chapter1_docx_import",
            "--source",
            str(_source(tmp_path)),
            "--apply",
            "--skip-existing",
        )

        from apps.registry.models import Operator

        assert Operator.objects.filter(employee_id="RUT-178162667").exists()

    def test_the_serial_is_still_read_verbatim_from_the_manual(self, tmp_path, capsys):
        """El informe describe la fuente: ver el espacio es lo que delata que el
        problema está en el documento y no en el registro DGAC."""
        report = _report(_source(tmp_path), capsys)

        assert report["aircraft"][0]["serial_number"] == SERIAL_DEL_MANUAL

    def test_it_is_stored_normalized(self, tmp_path):
        """Al crear, `Aircraft.save()` normaliza como siempre — esta fila no
        cambia eso, sólo la comparación."""
        call_command(
            "chapter1_docx_import", "--source", str(_source(tmp_path)), "--apply"
        )

        assert Aircraft.objects.get(registration="RPA-4401").serial_number == (
            normalize_serial(SERIAL_DEL_MANUAL)
        )


@pytest.mark.django_db
class TestWhatMustStillBeAConflict:
    def test_a_new_registration_on_an_existing_serial_is_still_a_conflict(
        self, tmp_path, capsys
    ):
        """El contrapeso, y la razón de que la segunda llave exista (`LV-134`):
        una matrícula que no está con una serie que sí **no es una aeronave
        nueva** — es la misma reinscrita o un dato mal transcrito, y crearla
        reventaría contra el índice único. Normalizar no puede volver esto
        invisible; al contrario, ahora lo detecta también con el espacio."""
        _aircraft("RPA-9999", SERIAL_EN_LA_BASE)

        report = _report(_source(tmp_path, registration="RPA-4401"), capsys)

        assert report["conflicts"]
        assert "RPA-4401" in report["conflicts"][0]
