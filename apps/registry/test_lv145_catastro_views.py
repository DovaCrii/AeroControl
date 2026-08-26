"""LV-145: the roster's four outputs, their permissions and their exports."""

from io import BytesIO

import pytest
from django.contrib.auth.models import Permission, User
from django.test import Client
from django.urls import reverse

from .catastro import aircraft_rows, build_catastro
from .models import Aircraft, CostCenter, Operator

ROUTES = [
    "registry-roster",
    "registry-roster-pdf",
    "registry-roster-xlsx",
    "registry-roster-csv",
]


@pytest.fixture
def world(db):
    center = CostCenter.objects.create(code="CC738", name="Los Pelambres")
    Aircraft.objects.create(
        registration="RPA-7126",
        type="Multirotor",
        # A leading "=" is what a spreadsheet executes on open. It goes in a
        # real field, not a synthetic one, because the export has to neutralize
        # every cell and not the ones someone remembered.
        model="=cmd|' /C calc'!A0",
        manufacturer="DJI",
        cost_center=center,
    )
    Operator.objects.create(
        full_name="Ana Rivas",
        employee_id="E-001",
        operator_type="+piloto",
        cost_center=center,
    )
    return {"cost_center": center}


def _client(username, *codenames):
    user = User.objects.create_user(username, password="password")
    for codename in codenames:
        user.user_permissions.add(Permission.objects.get(codename=codename))
    client = Client()
    assert client.login(username=username, password="password")
    return client


@pytest.fixture
def reader(world):
    """Someone with both halves -- the only shape that may read this report."""
    return _client("reader", "view_aircraft", "view_operator")


# -- who may read it -------------------------------------------------------


@pytest.mark.parametrize("route", ROUTES)
def test_an_anonymous_visitor_is_sent_to_the_login_page(world, route):
    response = Client().get(reverse(route))

    assert response.status_code == 302
    assert "/login" in response["Location"]


@pytest.mark.parametrize("route", ROUTES)
@pytest.mark.parametrize("codename", ["view_aircraft", "view_operator"])
def test_half_of_the_permission_is_not_enough(world, route, codename):
    """The document carries a fleet table **and** a personnel table, so holding
    one view permission must not hand over the other half. A single derived
    permission -- what `ModelPermissionRequiredMixin` does by default -- would
    have let `view_aircraft` alone print the whole roster of people."""
    client = _client("half", codename)

    assert client.get(reverse(route)).status_code == 403


@pytest.mark.parametrize("route", ROUTES)
def test_both_permissions_together_open_it(reader, route):
    assert reader.get(reverse(route)).status_code == 200


def test_the_menu_entry_needs_both_permissions_too(world):
    """A link that ends in a 403 teaches people to distrust the screen
    (LV-130), so the sidebar is gated on the same pair the view checks."""
    both = _client("both", "view_aircraft", "view_operator")
    only_fleet = _client("fleet", "view_aircraft")

    assert (
        reverse("registry-roster")
        in both.get(reverse("aircraft-list")).content.decode()
    )
    assert (
        reverse("registry-roster")
        not in only_fleet.get(reverse("aircraft-list")).content.decode()
    )


# -- the outputs -----------------------------------------------------------


@pytest.mark.parametrize(
    "route,content_type",
    [
        ("registry-roster-pdf", "application/pdf"),
        (
            "registry-roster-xlsx",
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        ),
        ("registry-roster-csv", "text/csv"),
    ],
)
def test_every_export_is_an_attachment_named_for_the_report(
    reader, route, content_type
):
    response = reader.get(reverse(route))

    assert content_type in response["Content-Type"]
    assert "aerocontrol-catastro" in response["Content-Disposition"]
    assert "attachment" in response["Content-Disposition"]


def test_the_pdf_wears_the_letterhead(reader):
    response = reader.get(reverse("registry-roster-pdf"))

    assert response.content[:5] == b"%PDF-"
    # Metadata is the only part of a reportlab PDF that is not compressed.
    assert b"J.E.J. Ingenier" in response.content
    assert b"Ref. CAT-" in response.content


def test_the_screen_shows_the_same_cells_the_pdf_prints(reader, world):
    """One `aircraft_rows` for the screen and the paper. The alternative -- the
    template reading the model directly -- is how a column ends up formatted
    one way on screen and another on the document someone files."""
    root = User.objects.create_superuser("root", "r@test.com", "password")
    response = reader.get(reverse("registry-roster"))

    on_screen = [row for _obj, row in response.context["aircraft_table"]]
    assert on_screen == aircraft_rows(build_catastro(root))


def test_the_cut_off_date_is_declared_on_the_screen_and_in_the_csv(reader):
    from django.utils import timezone

    today = timezone.localdate().isoformat()

    assert today in reader.get(reverse("registry-roster")).content.decode()
    assert today in reader.get(reverse("registry-roster-csv")).content.decode("utf-8")


# -- exports are inert -----------------------------------------------------


def test_a_formula_is_neutralised_in_the_csv(reader):
    body = reader.get(reverse("registry-roster-csv")).content.decode("utf-8")

    # Prefixed with an apostrophe, so the text stays readable and inert.
    assert "'=cmd" in body
    assert ",=cmd" not in body
    assert "'+piloto" in body


def test_a_formula_is_neutralised_in_the_xlsx(reader):
    from openpyxl import load_workbook

    workbook = load_workbook(
        BytesIO(reader.get(reverse("registry-roster-xlsx")).content)
    )

    values = [
        cell.value
        for sheet in workbook.worksheets
        for row in sheet.iter_rows()
        for cell in row
        if isinstance(cell.value, str)
    ]
    dangerous = [value for value in values if value.startswith(("=", "+", "-", "@"))]
    assert not dangerous, f"cells a spreadsheet would execute: {dangerous}"
    assert "'=cmd|' /C calc'!A0" in values


def test_the_xlsx_has_a_sheet_per_table_with_the_filter_on_the_header_row(reader):
    """The compliance report can hand openpyxl `sheet.dimensions` because its
    headers are on row 1. Here the cut-off sentences come first, so an
    auto-filter over the dimensions would take a sentence for a column title.
    """
    from openpyxl import load_workbook

    workbook = load_workbook(
        BytesIO(reader.get(reverse("registry-roster-xlsx")).content)
    )

    assert workbook.sheetnames == ["Flota", "Personal"]
    for sheet in workbook.worksheets:
        header_row = int(sheet.auto_filter.ref.split(":")[0].lstrip("A") or 1)
        assert header_row > 1, "the filter is anchored on the preamble"
        assert sheet.freeze_panes == f"A{header_row + 1}"
        assert sheet.cell(row=header_row, column=1).value


# -- query strings that should not break anything --------------------------


def test_a_malformed_cost_center_is_no_filter_and_not_a_500(reader):
    response = reader.get(reverse("registry-roster"), {"cost_center": "not-a-uuid"})

    assert response.status_code == 200
    assert response.context["catastro"]["totals"]["aircraft"] == 1


def test_an_unknown_status_is_no_filter_and_not_an_empty_roster(reader):
    """A stale bookmark should not answer "you have no aircraft"."""
    response = reader.get(reverse("registry-roster"), {"status": "teleported"})

    assert response.status_code == 200
    assert response.context["catastro"]["filters"].status == ""
    assert response.context["catastro"]["totals"]["aircraft"] == 1


def test_a_cost_center_filter_survives_into_the_export_links(reader, world):
    response = reader.get(
        reverse("registry-roster"), {"cost_center": str(world["cost_center"].pk)}
    )

    body = response.content.decode()
    assert f"cost_center={world['cost_center'].pk}" in body
