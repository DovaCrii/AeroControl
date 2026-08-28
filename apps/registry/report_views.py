"""Fleet and personnel roster: on-screen plus PDF / XLSX / CSV (LV-145).

Same shape as the compliance report next door -- a mixin that resolves the
filters from the query string, one view per output, and every number read from
a single data module (`catastro.py`) so the four cannot disagree.

The PDF wears the corporate letterhead from LV-144; this is the report that
letterhead was built shared for.
"""

import csv
from io import BytesIO

from django.http import HttpResponse
from django.utils.translation import gettext as _
from django.views.generic import TemplateView, View

from apps.compliance.digest import bucket_for
from apps.core.exports import neutralize
from apps.core.views import ModelPermissionRequiredMixin, lookup_by_pk
from .catastro import (
    AIRCRAFT_EXPIRY_COLUMN,
    AIRCRAFT_HEADERS,
    OPERATOR_EXPIRY_COLUMN,
    OPERATOR_HEADERS,
    CatastroFilters,
    aircraft_expiries,
    aircraft_rows,
    build_catastro,
    closing_total,
    operator_expiries,
    operator_rows,
    totals_sentence,
)
from .models import Aircraft, CostCenter

FILENAME_STEM = "aerocontrol-catastro"

# LV-167: carta **horizontal**. Con la columna de vigencia la flota pasa a nueve
# columnas, y en vertical (525.6 pt útiles) la matrícula y el número de serie se
# cortaban -- justo los dos datos con los que se identifica una aeronave. En
# horizontal hay 705.6, y los anchos se reparten a mano en vez de dejarlos a las
# partes iguales de reportlab: "Año" necesita 34 pt y "Modelo" ciento doce, y con
# columnas iguales la página se gasta en las cortas y las largas cortan a tres
# líneas.
AIRCRAFT_COL_WIDTHS = [72, 58, 112, 74, 112, 34, 66, 60, 117.6]
OPERATOR_COL_WIDTHS = [190, 100, 100, 130, 85, 100.6]


class CatastroMixin(ModelPermissionRequiredMixin):
    """Requires **both** view permissions, because the report has two halves.

    `ModelPermissionRequiredMixin` derives a single permission from `model` and
    `permission_action`; this document carries a fleet table and a personnel
    table, and someone holding `view_aircraft` without `view_operator` must not
    be handed the personnel one. `PermissionRequiredMixin` requires *all* of
    what `get_permission_required` returns, and the mixin already brings the
    right `handle_no_permission` (anonymous to the login page, authenticated
    without the permission to a 403).
    """

    # Kept for the mixin's contract and for anyone reading the class; the
    # permission tuple below is what actually gates the view.
    model = Aircraft
    permission_action = "view"

    def get_permission_required(self):
        return ("registry.view_aircraft", "registry.view_operator")

    def filters_for(self, request):
        cost_center = None
        if request.GET.get("cost_center"):
            cost_center = lookup_by_pk(
                CostCenter.objects.filter(is_active=True), request.GET["cost_center"]
            )
        # An unknown status is "no status filter", not an empty roster: the value
        # comes from a query string, and a stale bookmark should not answer
        # "you have no aircraft".
        status = request.GET.get("status", "")
        if status not in {code for code, _label in Aircraft.STATUS_CHOICES}:
            status = ""
        return CatastroFilters(
            cost_center=cost_center,
            status=status,
            include_terminal=request.GET.get("include_terminal") == "1",
            include_archived=request.GET.get("include_archived") == "1",
        )

    def catastro_for(self, request):
        return build_catastro(request.user, self.filters_for(request))


class CatastroReportView(CatastroMixin, TemplateView):
    template_name = "registry/catastro.html"

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        catastro = self.catastro_for(self.request)
        context.update(
            catastro=catastro,
            # Object and row side by side: the cells come from the same
            # `*_rows` the PDF prints, so the screen and the paper cannot drift
            # apart, while the object still supplies the link to the ficha.
            aircraft_table=list(zip(catastro["aircraft"], aircraft_rows(catastro))),
            operator_table=list(zip(catastro["operators"], operator_rows(catastro))),
            aircraft_headers=AIRCRAFT_HEADERS,
            operator_headers=OPERATOR_HEADERS,
            sentences=totals_sentence(catastro),
            closing=closing_total(catastro),
            title=_("Fleet and personnel roster"),
            cost_centers=CostCenter.objects.filter(is_active=True).order_by("code"),
            statuses=Aircraft.STATUS_CHOICES,
            filter_params=self.request.GET,
            export_query=self.request.GET.urlencode(),
        )
        return context


class CatastroReportPdfView(CatastroMixin, View):
    def get(self, request):
        from django.utils.html import escape
        from reportlab.platypus import PageBreak, Paragraph, Spacer, Table

        from apps.core import pdf as corepdf

        catastro = self.catastro_for(request)
        styles = corepdf.executive_stylesheet()
        title = str(_("AeroControl — Fleet and personnel roster"))
        letterhead = corepdf.Letterhead(
            title=title,
            reference=corepdf.derived_reference("CAT", catastro["as_of"]),
            generated_on=catastro["as_of"],
        )

        def urgency(expiries, column):
            """El color de urgencia de la celda de vigencia, fila por fila.

            LV-167: **la misma escala que el panel y el digest** -- una fecha
            ámbar en pantalla es ámbar en el papel. El tramo lo calcula acá con
            `bucket_for`, que es el dueño de los cortes, y `catastro.py` sólo
            entrega las fechas: que el padrón importara `apps.compliance` para
            colorear una celda sería atarlo al cumplimiento por el color.

            Un nulo **no entra en ningún tramo** y por lo tanto no se pinta: es
            la lección de `LV-29`, un nulo es "nunca se ingresó" y no "vigente".

            Y acá se ve por qué `urgency_commands` toma **rangos**: esto pinta
            una celda por fila (cada fila tiene su fecha y su tramo), mientras el
            informe de cumplimiento pinta columnas enteras. Cualquiera de las dos
            orientaciones fijada en el helper habría tenido que reescribirse.
            """
            today = catastro["as_of"]
            return corepdf.urgency_commands(
                (bucket_for(expiry, today), (column, index), (column, index))
                for index, expiry in enumerate(expiries, start=1)
                if expiry is not None
            )

        def table(headers, rows, widths, empty_message, extra=()):
            if not rows:
                return Paragraph(escape(str(empty_message)), styles["Normal"])
            # Every cell is a Paragraph, headers included: at these widths
            # "Manufacturer" has to wrap, and a plain string in a reportlab
            # table does not wrap -- it overruns into the next column.
            # `escape` because Paragraph parses its text as mini-HTML, and an
            # "&" in a manufacturer name would raise instead of printing.
            data = [
                [Paragraph(escape(str(head)), styles["CellHeader"]) for head in headers]
            ]
            data += [
                [
                    Paragraph(escape(str(neutralize(cell))), styles["Cell"])
                    for cell in row
                ]
                for row in rows
            ]
            grid = Table(data, colWidths=widths, repeatRows=1)
            grid.setStyle(corepdf.executive_table_style(extra))
            return grid

        elements = [Paragraph(escape(title), styles["Title"])]
        for sentence in totals_sentence(catastro):
            elements.append(Paragraph(escape(str(sentence)), styles["Normal"]))
        elements.append(Spacer(1, 14))

        # Escaped like every other Paragraph, catalog string or not: an "&" in a
        # translation is the same crash as an "&" in a manufacturer name.
        elements.append(Paragraph(escape(str(_("Fleet"))), styles["Heading2"]))
        elements.append(
            table(
                AIRCRAFT_HEADERS,
                aircraft_rows(catastro),
                AIRCRAFT_COL_WIDTHS,
                _("No aircraft registered."),
                urgency(aircraft_expiries(catastro), AIRCRAFT_EXPIRY_COLUMN),
            )
        )
        # LV-164 (pedido del usuario, 2026-08-27): el personal arranca en hoja
        # nueva. Son dos padrones de cosas distintas, y en un documento que se
        # entrega cada tabla se lee —y se fotocopia, y se firma— por separado.
        # El `keepWithNext` del estilo evita además que cualquiera de los dos
        # títulos quede huérfano al pie; el salto es la decisión editorial, no
        # el parche.
        elements.append(PageBreak())
        elements.append(Paragraph(escape(str(_("Personnel"))), styles["Heading2"]))
        elements.append(
            table(
                OPERATOR_HEADERS,
                operator_rows(catastro),
                OPERATOR_COL_WIDTHS,
                _("No operators registered."),
                urgency(operator_expiries(catastro), OPERATOR_EXPIRY_COLUMN),
            )
        )

        # LV-172: el total de cierre, al final del documento y no sólo en el
        # encabezado. Con la flota y el personal ocupando varias páginas, quien
        # recibe el catastro termina de leer a varias páginas del número que lo
        # resume -- y es al final donde se comprueba que no se cortó nada. Va en
        # negrita, que es lo que lo separa de las notas al pie de cada tabla.
        elements.append(Spacer(1, 12))
        elements.append(
            Paragraph(
                f"<b>{escape(str(closing_total(catastro)))}</b>", styles["Normal"]
            )
        )

        # LV-167: horizontal. Ver el comentario de los anchos: con nueve columnas
        # la vertical cortaba la matrícula y el número de serie.
        from reportlab.lib.pagesizes import landscape, letter

        return corepdf.pdf_response(
            elements, letterhead, f"{FILENAME_STEM}.pdf", pagesize=landscape(letter)
        )


class CatastroReportXlsxView(CatastroMixin, View):
    def get(self, request):
        from openpyxl import Workbook
        from openpyxl.styles import Font
        from openpyxl.utils import get_column_letter

        catastro = self.catastro_for(request)
        sentences = totals_sentence(catastro)
        workbook = Workbook()

        def fill(sheet, headers, rows):
            """Preamble, then the table, with the filter anchored on the header
            row rather than on `sheet.dimensions`.

            The compliance report puts its headers on row 1 and can hand
            openpyxl the whole sheet; here the cut-off sentences come first, so
            an auto-filter over the dimensions would take the sentences for
            column titles.

            Written by coordinate and not with `append`: an `append([])` for the
            blank separator advances openpyxl's internal cursor but creates no
            cell, so `max_row` still points at the sentence above it -- which
            put the bold font on an empty row and the filter one row short of
            the headers.
            """
            for offset, sentence in enumerate(sentences, start=1):
                sheet.cell(row=offset, column=1, value=str(sentence))
            header_row = len(sentences) + 2  # one blank row of separation
            for column, head in enumerate(headers, start=1):
                sheet.cell(row=header_row, column=column, value=str(head)).font = Font(
                    bold=True
                )
            for index, row in enumerate(rows, start=header_row + 1):
                for column, value in enumerate(row, start=1):
                    sheet.cell(row=index, column=column, value=neutralize(value))
            last_row = header_row + len(rows)
            sheet.freeze_panes = f"A{header_row + 1}"
            sheet.auto_filter.ref = (
                f"A{header_row}:{get_column_letter(len(headers))}{last_row}"
            )
            for column in sheet.columns:
                sheet.column_dimensions[column[0].column_letter].width = min(
                    max(len(str(cell.value or "")) for cell in column) + 2, 40
                )

        fleet = workbook.active
        fleet.title = "Flota"
        fill(fleet, AIRCRAFT_HEADERS, aircraft_rows(catastro))
        fill(
            workbook.create_sheet("Personal"),
            OPERATOR_HEADERS,
            operator_rows(catastro),
        )

        output = BytesIO()
        workbook.save(output)
        response = HttpResponse(
            output.getvalue(),
            content_type=(
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
            ),
        )
        response["Content-Disposition"] = f'attachment; filename="{FILENAME_STEM}.xlsx"'
        return response


class CatastroReportCsvView(CatastroMixin, View):
    def get(self, request):
        catastro = self.catastro_for(request)
        response = HttpResponse(content_type="text/csv; charset=utf-8")
        response["Content-Disposition"] = f'attachment; filename="{FILENAME_STEM}.csv"'
        response.write("﻿")  # BOM so Excel detects UTF-8
        writer = csv.writer(response, lineterminator="\r\n")
        writer.writerow([_("Fleet and personnel roster")])
        for sentence in totals_sentence(catastro):
            writer.writerow([neutralize(sentence)])
        writer.writerow([])

        for heading, headers, rows in (
            (_("Fleet"), AIRCRAFT_HEADERS, aircraft_rows(catastro)),
            (_("Personnel"), OPERATOR_HEADERS, operator_rows(catastro)),
        ):
            writer.writerow([neutralize(heading)])
            writer.writerow([str(head) for head in headers])
            for row in rows:
                writer.writerow([neutralize(value) for value in row])
            writer.writerow([])
        return response
